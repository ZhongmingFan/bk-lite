from io import BytesIO
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from apps.node_mgmt.constants.node import NodeConstants
from apps.system_mgmt.models import Group

EXPORT_LIMIT = 5000
COLLECTOR_JOIN = "；"
ORG_JOIN = ","

EXPORT_HEADERS_ZH = [
    "节点名称",
    "IP",
    "操作系统",
    "CPU 架构",
    "节点类型",
    "安装方式",
    "组织",
    "在线状态",
    "最后上报时间",
    "控制器版本",
    "控制器可升级",
    "托管组件",
]

EXPORT_HEADERS_EN = [
    "Node Name",
    "IP",
    "OS",
    "CPU Architecture",
    "Node Type",
    "Install Method",
    "Organization",
    "Online Status",
    "Last Report Time",
    "Controller Version",
    "Controller Upgradeable",
    "Hosted Collectors",
]


def is_english_locale(locale):
    return str(locale or "").lower().startswith("en")


def export_labels(locale):
    english = is_english_locale(locale)
    if english:
        return {
            "unassigned": "Unassigned",
            "online": "Online",
            "offline": "Offline",
            "yes": "Yes",
            "no": "No",
            "os": {"linux": "Linux", "windows": "Windows"},
            "install_method": {"auto": "Remote", "manual": "Manual"},
            "node_type": {"host": "Host", "container": "Container"},
            "collector_status": {
                0: "Normal",
                1: "Unknown",
                2: "Error",
                3: "Stopped",
                4: "Not Started",
                10: "Installing",
                11: "Not Started",
                12: "Failed to Install",
            },
            "headers": EXPORT_HEADERS_EN,
            "empty": "No nodes to export",
            "over_limit": "Export exceeds the limit of {limit} nodes, current count is {count}.",
        }
    return {
        "unassigned": "未归属",
        "online": "在线",
        "offline": "离线",
        "yes": "是",
        "no": "否",
        "os": {NodeConstants.LINUX_OS: NodeConstants.LINUX_OS_DISPLAY, NodeConstants.WINDOWS_OS: NodeConstants.WINDOWS_OS_DISPLAY},
        "install_method": {"auto": "远程", "manual": "手动"},
        "node_type": {"host": "主机节点", "container": "容器节点"},
        "collector_status": {
            0: "正常",
            1: "未知的",
            2: "异常",
            3: "停止",
            4: "未启动",
            10: "安装中",
            11: "未启动",
            12: "安装失败",
        },
        "headers": EXPORT_HEADERS_ZH,
        "empty": "没有可导出的节点",
        "over_limit": "导出数量超过上限 {limit}，当前 {count} 台",
    }


def format_organization_cell(organization_ids, organization_names, unassigned_label):
    names = []
    for org_id in organization_ids or []:
        try:
            key = int(org_id)
        except (TypeError, ValueError):
            continue
        name = organization_names.get(key)
        if name:
            names.append(name)
    return ORG_JOIN.join(names) if names else unassigned_label


def _hosted_collectors(node):
    status = node.get("status") if isinstance(node.get("status"), dict) else {}
    reported = status.get("collectors") or []
    installed = status.get("collectors_install") or []
    if not isinstance(reported, list):
        reported = []
    if not isinstance(installed, list):
        installed = []
    seen = {item.get("collector_id") for item in reported if isinstance(item, dict)}
    extra = [item for item in installed if isinstance(item, dict) and item.get("collector_id") not in seen]
    return [item for item in reported if isinstance(item, dict)] + extra


def format_hosted_collectors_cell(collectors, status_labels):
    parts = []
    for item in collectors or []:
        name = str(item.get("collector_name") or item.get("name") or "").strip()
        if not name:
            continue
        try:
            code = int(item.get("status"))
        except (TypeError, ValueError):
            code = 1
        label = status_labels.get(code, status_labels.get(1, ""))
        parts.append(f"{name}:{label}")
    return COLLECTOR_JOIN.join(parts)


def _controller_version(node):
    for item in node.get("versions") or []:
        if item.get("component_type") == "controller":
            version = item.get("version") or ""
            upgradeable = bool(item.get("upgradeable"))
            return version, upgradeable
    return "", False


def build_export_row(node, labels, organization_names):
    os_value = node.get("operating_system") or ""
    install_method = node.get("install_method") or ""
    node_type = node.get("node_type") or ""
    version, upgradeable = _controller_version(node)
    return [
        node.get("name") or "",
        node.get("ip") or "",
        labels["os"].get(os_value, os_value),
        node.get("cpu_architecture") or "",
        labels["node_type"].get(node_type, node_type),
        labels["install_method"].get(install_method, install_method),
        format_organization_cell(node.get("organization") or [], organization_names, labels["unassigned"]),
        labels["online"] if node.get("active") else labels["offline"],
        node.get("updated_at") or "",
        version,
        labels["yes"] if upgradeable else labels["no"],
        format_hosted_collectors_cell(_hosted_collectors(node), labels["collector_status"]),
    ]


def load_organization_names(organization_ids):
    ids = []
    for org_id in organization_ids:
        try:
            ids.append(int(org_id))
        except (TypeError, ValueError):
            continue
    if not ids:
        return {}
    return dict(Group.objects.filter(id__in=ids).values_list("id", "name"))


def build_export_filename(region_name, stamp):
    safe_region = str(region_name or "cloud").replace("/", "_")
    return f"节点清单_{safe_region}_{stamp}.xlsx"


def build_export_workbook_bytes(rows, headers):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "nodes"
    sheet.append(headers)
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    header_alignment = Alignment(horizontal="center", vertical="center")
    for col_num, _ in enumerate(headers, 1):
        cell = sheet.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = header_alignment
    for row in rows:
        sheet.append(row)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def content_disposition(filename):
    ascii_name = "nodes.xlsx"
    return f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(filename)}"
