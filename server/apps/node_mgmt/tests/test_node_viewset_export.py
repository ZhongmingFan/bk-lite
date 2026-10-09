"""NodeViewSet.export_excel：筛选导出、已选 ID、空结果、超限。"""
import json
import uuid
from datetime import timedelta
from io import BytesIO
from unittest.mock import patch

import pytest
from django.utils import timezone as dj_timezone
from openpyxl import load_workbook
from rest_framework.test import APIRequestFactory, force_authenticate

from apps.base.tests.factories import UserFactory
from apps.node_mgmt.constants.node import NodeConstants
from apps.node_mgmt.models import CloudRegion, Collector, Node
from apps.node_mgmt.models.sidecar import NodeOrganization
from apps.node_mgmt.views import node as node_view
from apps.node_mgmt.views.node import NodeViewSet
from apps.system_mgmt.models import Group

pytestmark = pytest.mark.django_db
factory = APIRequestFactory()


def _user():
    user = UserFactory(username=f"node-ex-{uuid.uuid4().hex[:8]}", domain="domain.com", is_superuser=True)
    user.permission = {"node": {"cloud_region_node-View"}}
    user.locale = "zh-Hans"
    return user


def _auth(request, user=None):
    user = user or _user()
    force_authenticate(request, user=user)
    request.COOKIES["current_team"] = "1"
    return user


def _region_and_nodes():
    region = CloudRegion.objects.create(name="export-region")
    keep = Node.objects.create(
        id=f"keep-{uuid.uuid4().hex[:8]}",
        name="keep-node",
        ip="10.0.0.1",
        operating_system=NodeConstants.LINUX_OS,
        collector_configuration_directory="/tmp",
        cloud_region=region,
        status={"collectors": [{"collector_id": "telegraf_linux", "status": 2}]},
    )
    other = Node.objects.create(
        id=f"other-{uuid.uuid4().hex[:8]}",
        name="other-node",
        ip="10.0.0.2",
        operating_system=NodeConstants.LINUX_OS,
        collector_configuration_directory="/tmp",
        cloud_region=region,
        status={},
    )
    Node.objects.filter(id=keep.id).update(updated_at=dj_timezone.now() - timedelta(seconds=10))
    Node.objects.filter(id=other.id).update(updated_at=dj_timezone.now() - timedelta(seconds=120))
    NodeOrganization.objects.create(node=keep, organization=7)
    NodeOrganization.objects.create(node=other, organization=8)
    Group.objects.get_or_create(id=7, defaults={"name": "alpha", "parent_id": 0})
    Group.objects.get_or_create(id=8, defaults={"name": "beta", "parent_id": 0})
    Collector.objects.create(
        id="telegraf_linux",
        name="Telegraf",
        service_type="exec",
        node_operating_system="linux",
        executable_path="/bin/telegraf",
        execute_parameters="",
        created_by="tester",
        updated_by="tester",
    )
    return region, keep, other


def _post_export(monkeypatch, nodes, body, query=""):
    ids = [node.id for node in nodes]
    monkeypatch.setattr(node_view, "get_node_permission", lambda request: {"team": [1], "instance": []})
    monkeypatch.setattr(
        node_view,
        "get_catalog_node_queryset",
        lambda request, permission=None: Node.objects.filter(id__in=ids),
    )
    request = factory.post(f"/node/export_excel/{query}", body, format="json")
    _auth(request)
    return NodeViewSet.as_view({"post": "export_excel"})(request)


def test_export_filtered_offline_nodes(monkeypatch):
    region, keep, other = _region_and_nodes()
    resp = _post_export(
        monkeypatch,
        [keep, other],
        {
            "cloud_region_id": region.id,
            "filters": {"active": [{"lookup_expr": "in", "value": ["false"]}]},
        },
    )
    assert resp.status_code == 200
    sheet = load_workbook(BytesIO(resp.content)).active
    names = [row[0].value for row in sheet.iter_rows(min_row=2)]
    assert names == ["other-node"]
    assert "other-node" in names and "keep-node" not in names


def test_export_selected_ids_ignores_filters(monkeypatch):
    region, keep, other = _region_and_nodes()
    resp = _post_export(
        monkeypatch,
        [keep, other],
        {
            "cloud_region_id": region.id,
            "selected_ids": [keep.id, other.id],
            "filters": {"active": [{"lookup_expr": "in", "value": ["false"]}]},
        },
    )
    assert resp.status_code == 200
    names = [row[0].value for row in load_workbook(BytesIO(resp.content)).active.iter_rows(min_row=2)]
    assert set(names) == {"keep-node", "other-node"}


def test_export_drops_unknown_selected_id(monkeypatch):
    region, keep, other = _region_and_nodes()
    resp = _post_export(
        monkeypatch,
        [keep, other],
        {"cloud_region_id": region.id, "selected_ids": [keep.id, "missing-node"]},
    )
    assert resp.status_code == 200
    names = [row[0].value for row in load_workbook(BytesIO(resp.content)).active.iter_rows(min_row=2)]
    assert names == ["keep-node"]


def test_export_empty_selected_fails_without_file(monkeypatch):
    region, keep, other = _region_and_nodes()
    resp = _post_export(
        monkeypatch,
        [keep, other],
        {"cloud_region_id": region.id, "selected_ids": ["missing-node"]},
    )
    assert resp.status_code == 400
    assert resp["Content-Type"].startswith("application/json")
    assert "没有可导出的节点" in json.loads(resp.content)["message"]


def test_export_empty_filter_fails_without_file(monkeypatch):
    region, keep, other = _region_and_nodes()
    resp = _post_export(
        monkeypatch,
        [keep, other],
        {
            "cloud_region_id": region.id,
            "filters": {"name": [{"lookup_expr": "icontains", "value": "no-such-node"}]},
        },
    )
    assert resp.status_code == 400
    assert "没有可导出的节点" in json.loads(resp.content)["message"]


def test_export_over_limit_fails_without_truncation(monkeypatch):
    region, keep, other = _region_and_nodes()
    with patch("apps.node_mgmt.views.node.EXPORT_LIMIT", 1):
        resp = _post_export(monkeypatch, [keep, other], {"cloud_region_id": region.id})
    assert resp.status_code == 400
    assert "上限" in json.loads(resp.content)["message"]
    assert resp.get("Content-Disposition") is None


def test_export_selected_ids_not_list_returns_400(monkeypatch):
    region, keep, other = _region_and_nodes()
    resp = _post_export(
        monkeypatch,
        [keep, other],
        {"cloud_region_id": region.id, "selected_ids": "keep-id"},
    )
    assert resp.status_code == 400
    assert resp.get("Content-Disposition") is None


def test_export_hosted_cell_uses_collector_display_name(monkeypatch):
    region, keep, other = _region_and_nodes()
    resp = _post_export(
        monkeypatch,
        [keep],
        {"cloud_region_id": region.id, "selected_ids": [keep.id]},
    )
    sheet = load_workbook(BytesIO(resp.content)).active
    hosted = sheet[2][11].value
    assert hosted == "Telegraf:异常"
